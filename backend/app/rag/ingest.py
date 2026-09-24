from pathlib import Path
import hashlib
import io
from typing import List, Dict, Any, Generator
from sqlalchemy.orm import Session
from app.db.models import Document, Chunk, DocumentStatus, EquipmentClass
from app.db.session import SessionLocal
from app.core.config import get_settings
from app.core.paths import RAW_DOCS_DIR

settings = get_settings()

# Heavy optional parsers (PDF/DOCX/OCR). Text parsing always works.
try:
    import fitz  # PyMuPDF
    _PDF_OK = True
except Exception:
    fitz = None
    _PDF_OK = False

try:
    import pdfplumber
    _PDFPLUMBER_OK = True
except Exception:
    pdfplumber = None
    _PDFPLUMBER_OK = False

try:
    from docx import Document as DocxDocument
    _DOCX_OK = True
except Exception:
    DocxDocument = None
    _DOCX_OK = False

try:
    import pytesseract
    from PIL import Image
    _OCR_OK = True
except Exception:
    pytesseract = None
    Image = None
    _OCR_OK = False


class DocumentParser:
    def parse(self, file_path: str, doc_id: int) -> Generator[Dict[str, Any], None, None]:
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            yield from self._parse_pdf(file_path, doc_id)
        elif ext == ".docx":
            yield from self._parse_docx(file_path, doc_id)
        elif ext == ".txt":
            yield from self._parse_txt(file_path, doc_id)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    
    def _parse_pdf(self, file_path: str, doc_id: int) -> Generator[Dict[str, Any], None, None]:
        if not _PDF_OK:
            raise ValueError("PDF parsing requires pymupdf (pip install -r requirements-heavy.txt)")
        doc = fitz.open(file_path)
        page_count = len(doc)
        
        for page_num in range(page_count):
            page = doc[page_num]
            text = page.get_text("text")
            
            tables = []
            if _PDFPLUMBER_OK:
                try:
                    with pdfplumber.open(file_path) as pdf:
                        if page_num < len(pdf.pages):
                            pl_page = pdf.pages[page_num]
                            extracted_tables = pl_page.extract_tables()
                            if extracted_tables:
                                for table in extracted_tables:
                                    tables.append(self._table_to_markdown(table))
                except Exception:
                    pass
            
            if not text.strip() and _OCR_OK:
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text = pytesseract.image_to_string(img)
            
            if tables:
                text += "\n\nTables:\n" + "\n\n".join(tables)
            
            yield {
                "page_start": page_num + 1,
                "page_end": page_num + 1,
                "section_title": f"Page {page_num + 1}",
                "text": text.strip(),
                "doc_id": doc_id
            }
        doc.close()
    
    def _parse_docx(self, file_path: str, doc_id: int) -> Generator[Dict[str, Any], None, None]:
        if not _DOCX_OK:
            raise ValueError("DOCX parsing requires python-docx (pip install -r requirements-heavy.txt)")
        doc = DocxDocument(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row])
            full_text.append(self._table_to_markdown(table_data))
        
        text = "\n".join(full_text)
        yield {
            "page_start": 1,
            "page_end": 1,
            "section_title": "Document",
            "text": text,
            "doc_id": doc_id
        }
    
    def _parse_txt(self, file_path: str, doc_id: int) -> Generator[Dict[str, Any], None, None]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        yield {
            "page_start": 1,
            "page_end": 1,
            "section_title": "Document",
            "text": text,
            "doc_id": doc_id
        }
    
    def _table_to_markdown(self, table: List[List[str]]) -> str:
        if not table:
            return ""
        headers = table[0]
        rows = table[1:]
        md = "| " + " | ".join(headers) + " |\n"
        md += "| " + " | ".join(["---"] * len(headers)) + " |\n"
        for row in rows:
            md += "| " + " | ".join(row) + " |\n"
        return md


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    cleaned = []
    for line in lines:
        line = line.strip()
        if line:
            cleaned.append(line)
    return "\n".join(cleaned)


def chunk_text(text: str, max_tokens: int = 800, overlap: int = 100) -> List[str]:
    words = text.split()
    if len(words) <= max_tokens:
        return [text]
    
    chunks = []
    for i in range(0, len(words), max_tokens - overlap):
        chunk = " ".join(words[i:i + max_tokens])
        chunks.append(chunk)
        if i + max_tokens >= len(words):
            break
    return chunks


def detect_equipment_and_test(text: str, db: Session) -> tuple:
    equip_classes = db.query(EquipmentClass).all()
    equip_names = []
    for ec in equip_classes:
        equip_names.append(ec.name)
        equip_names.extend(ec.aliases or [])
    
    from rapidfuzz import process, fuzz
    match = process.extractOne(text[:500], equip_names, scorer=fuzz.WRatio, score_cutoff=75)
    equipment_class_id = None
    if match:
        for ec in equip_classes:
            if match[0] == ec.name or match[0] in (ec.aliases or []):
                equipment_class_id = ec.id
                break
    
    return equipment_class_id, None


def process_document(document_id: int):
    db = SessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == document_id).first()
        if not doc:
            return
        
        file_path = RAW_DOCS_DIR / doc.filename
        if not file_path.exists():
            doc.status = "failed"
            db.commit()
            return
        
        parser = DocumentParser()
        all_chunks = []
        page_count = 0
        
        for parsed in parser.parse(str(file_path), document_id):
            page_count = max(page_count, parsed["page_end"])
            cleaned = clean_text(parsed["text"])
            if not cleaned:
                continue
            
            equipment_class_id, test_id = detect_equipment_and_test(cleaned, db)
            
            chunks = chunk_text(cleaned)
            for idx, chunk_text_content in enumerate(chunks):
                all_chunks.append({
                    "document_id": document_id,
                    "page_start": parsed["page_start"],
                    "page_end": parsed["page_end"],
                    "section_title": parsed["section_title"],
                    "text": chunk_text_content,
                    "equipment_class_id": equipment_class_id,
                    "test_id": test_id,
                    "chunk_index": idx,
                    "token_count": len(chunk_text_content.split())
                })
        
        for chunk_data in all_chunks:
            chunk = Chunk(**chunk_data)
            db.add(chunk)
        db.commit()
        
        # Embed + index for retrieval (no-op when heavy deps are unavailable)
        try:
            from app.rag.retriever import retriever
            stored_chunks = db.query(Chunk).filter(Chunk.document_id == document_id).all()
            retriever.add_chunks(stored_chunks)
        except Exception:
            pass
        
        doc.page_count = page_count
        doc.status = DocumentStatus.READY
        db.commit()
        
    except Exception:
        doc = db.query(Document).filter(Document.id == document_id).first()
        if doc:
            doc.status = DocumentStatus.FAILED
            db.commit()
        raise
    finally:
        db.close()